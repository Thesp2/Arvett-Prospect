"""
Export — write scored leads to a CRM-ready CSV and a color-tiered XLSX.
Columns are defined by the vertical (vertical.columns), so the same exporter
serves every use case. Generalized from data-kit/enrich_leads.py::write_outputs.
"""
from __future__ import annotations

import csv
import json


def _cell_value(value):
    """Convert structured provenance into a readable scalar for CSV/XLSX."""
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if value is None:
        return ""
    return value


def _ordered_columns(columns):
    preferred = [
        "name", "phone", "business_email", "owner_name", "owner_role",
        "city", "city_normalized", "address", "category", "website_status",
        "website", "score", "contact_confidence_score", "opener", "why",
        "source_url", "contact_source", "checked_at",
    ]
    rank = {key: index for index, key in enumerate(preferred)}
    return sorted(enumerate(columns), key=lambda item: (rank.get(item[1][1], 1000), item[0]))


def quality_warnings(record: dict) -> list[str]:
    """Return explicit pre-export warnings without rewriting source evidence."""
    warnings = []
    if not record.get("name"):
        warnings.append("missing business name")
    if record.get("lat") is not None or record.get("lon") is not None:
        try:
            if not (-90 <= float(record.get("lat")) <= 90 and -180 <= float(record.get("lon")) <= 180):
                warnings.append("invalid coordinates")
        except (TypeError, ValueError):
            warnings.append("invalid coordinates")
    if record.get("website_status") == "FUNCTIONAL_WEBSITE" and not record.get("audit"):
        warnings.append("functional status lacks audit evidence")
    if not record.get("source") and not record.get("source_references"):
        warnings.append("missing source provenance")
    for key in ("score", "website_opportunity_score", "contact_confidence_score"):
        try:
            if not 0 <= int(record.get(key, 0)) <= 100:
                warnings.append(f"{key} outside 0-100")
        except (TypeError, ValueError):
            warnings.append(f"{key} is not numeric")
    return warnings


def write_csv(leads: list[dict], columns: list[tuple[str, str]], path: str) -> str:
    columns = [column for _, column in _ordered_columns(columns)]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([h for h, _ in columns])
        for r in leads:
            w.writerow([_cell_value(r.get(k, "")) for _, k in columns])
    return path


def write_xlsx(leads: list[dict], columns: list[tuple[str, str]], path: str,
               title: str = "Qualified Leads") -> str:
    """Write a review-ready workbook with qualified, complete, and summary sheets."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter
    from collections import Counter

    wb = openpyxl.Workbook()
    HEAD_F = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    HEAD_FILL = PatternFill("solid", start_color="1F4E78")
    BODY = Font(name="Arial", size=10)
    WRAP = Alignment(wrap_text=True, vertical="top")
    TIER_FILL = {"A": PatternFill("solid", start_color="C6EFCE"),
                 "B": PatternFill("solid", start_color="FFEB9C"),
                 "C": PatternFill("solid", start_color="FFC7CE")}
    columns = [column for _, column in _ordered_columns(columns)]
    link_keys = {"website", "source_url", "contact_source", "website_audit_url"}

    def add_lead_sheet(ws, rows):
        for i, (header, _) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=i, value=header)
            cell.font = HEAD_F
            cell.fill = HEAD_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = max(12, min(38, len(header) + 5))
        tier_col = next((i for i, (_, key) in enumerate(columns, 1) if key == "tier"), None)
        for ri, record in enumerate(rows, 2):
            for ci, (_, key) in enumerate(columns, 1):
                cell = ws.cell(row=ri, column=ci, value=_cell_value(record.get(key, "")))
                cell.font = BODY
                cell.alignment = WRAP
                if key in link_keys and isinstance(record.get(key), str) and record.get(key).startswith(("http://", "https://")):
                    cell.hyperlink = record[key]
                    cell.style = "Hyperlink"
                if key in {"phone", "original_phone", "normalized_phone", "zip"}:
                    cell.number_format = "@"
                if ri % 2 == 0 and not (tier_col and record.get("tier") in TIER_FILL):
                    cell.fill = PatternFill("solid", start_color="F5F8FC")
            if tier_col and record.get("tier") in TIER_FILL:
                ws.cell(row=ri, column=tier_col).fill = TIER_FILL[record["tier"]]
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 32
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows)+1}"
        score_col = next((i for i, (_, key) in enumerate(columns, 1)
                          if key == "website_opportunity_score"), None)
        if score_col and rows:
            letter = get_column_letter(score_col)
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{len(rows) + 1}",
                ColorScaleRule(start_type="min", start_color="F8696B",
                               mid_type="percentile", mid_value=50, mid_color="FFEB84",
                               end_type="max", end_color="63BE7B"),
            )

    # Keep all rows accessible; qualified sheet is a prioritization aid, not deletion.
    qualified = [
        r for r in leads
        if int(r.get("score") or 0) >= 35 and not r.get("excluded_reason")
    ]
    ws = wb.active
    ws.title = (title or "Qualified Leads")[:31]
    add_lead_sheet(ws, qualified)

    all_ws = wb.create_sheet("All Businesses")
    add_lead_sheet(all_ws, leads)

    summary = wb.create_sheet("Search Summary")
    summary.append(["Metric", "Count / value"])
    for cell in summary[1]:
        cell.font = HEAD_F
        cell.fill = HEAD_FILL
    status_counts = Counter(str(r.get("website_status") or "Unknown") for r in leads)
    source_counts = Counter(str(r.get("source") or "Unknown") for r in leads)
    markets = sorted({str(r.get("target_market") or "") for r in leads if r.get("target_market")})
    summary_rows = [
        ("Target market", ", ".join(markets)),
        ("Businesses collected", len(leads)),
        ("Qualified leads (score >= 35; exclusions removed)", len(qualified)),
        ("Excluded records", sum(bool(r.get("excluded_reason")) for r in leads)),
        ("Records with quality warnings", sum(bool(quality_warnings(r)) for r in leads)),
        ("With public phone", sum(bool(r.get("phone")) for r in leads)),
        ("With business email", sum(bool(r.get("business_email") or r.get("email")) for r in leads)),
        ("Owner / decision-maker identified", sum(bool(r.get("owner_name")) for r in leads)),
        ("Location source-confirmed city", sum(bool(r.get("city")) for r in leads)),
        ("Website status — no URL found", status_counts.get("NO_WEBSITE_FOUND", 0)),
        ("Website status — social-only", status_counts.get("SOCIAL_ONLY", 0)),
        ("Website status — broken", status_counts.get("WEBSITE_BROKEN", 0)),
        ("Website status — outdated", status_counts.get("OUTDATED_WEBSITE", 0)),
    ]
    summary_rows.extend([
        ("Missing phone %", f"{(sum(not bool(r.get('phone')) for r in leads) / len(leads) * 100):.1f}%"
         if leads else "0.0%"),
        ("Missing email %", f"{(sum(not bool(r.get('business_email') or r.get('email')) for r in leads) / len(leads) * 100):.1f}%"
         if leads else "0.0%"),
    ])
    summary_rows.extend((f"Source — {key}", value) for key, value in sorted(source_counts.items()))
    summary_rows.extend((f"Website status — {key}", value) for key, value in sorted(status_counts.items()))
    for row in summary_rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 48
    summary.column_dimensions["B"].width = 26
    for row in summary.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY
            cell.alignment = WRAP
    summary.freeze_panes = "A2"
    wb.save(path)
    return path


def write_outputs(leads, columns, stem: str, log=print) -> tuple[str, str]:
    csv_path = write_csv(leads, columns, f"{stem}_crm.csv")
    log(f"  -> {csv_path}")
    xlsx_path = write_xlsx(leads, columns, f"{stem}.xlsx", title="Qualified Leads")
    log(f"  -> {xlsx_path}")
    return csv_path, xlsx_path
