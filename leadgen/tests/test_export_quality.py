import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from openpyxl import load_workbook

from leadgen.export import write_xlsx
from leadgen.verticals.web_design import COLUMNS


def test_contact_first_workbook_has_links_and_filters():
    row = {
        "name": "Example Business",
        "phone": "",
        "business_email": "",
        "website": "https://example.test",
        "source_url": "https://openstreetmap.org/node/1",
        "website_status": "FUNCTIONAL_WEBSITE",
        "website_opportunity_score": 40,
        "contact_confidence_score": 20,
        "score": 40,
        "tier": "B",
        "audit": {"reachable": True},
    }
    with tempfile.TemporaryDirectory() as directory:
        path = os.path.join(directory, "report.xlsx")
        write_xlsx([row], COLUMNS, path)
        workbook = load_workbook(path)
        assert workbook.sheetnames == ["Qualified Leads", "All Businesses", "Search Summary"]
        sheet = workbook["All Businesses"]
        headers = [cell.value for cell in sheet[1]]
        assert headers.index("Business Name") < headers.index("Opportunity score")
        website_cell = sheet.cell(row=2, column=headers.index("Website URL") + 1)
        assert website_cell.hyperlink.target == "https://example.test"
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref
        assert workbook["Search Summary"]["A1"].value == "Metric"
