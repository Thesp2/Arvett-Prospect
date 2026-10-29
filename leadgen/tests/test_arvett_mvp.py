import os
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from leadgen.pipeline import _matches_business_category
from leadgen.verticals.web_design import _score
from leadgen.export import write_xlsx


class ArvettMVPTests(unittest.TestCase):
    def test_category_groups(self):
        self.assertTrue(_matches_business_category({"category": "cafe"}, "food"))
        self.assertTrue(_matches_business_category({"category": "furniture"}, "retail"))
        self.assertTrue(_matches_business_category({"category": "clinic"}, "healthcare"))
        self.assertFalse(_matches_business_category({"category": "fuel"}, "food"))
        self.assertTrue(_matches_business_category({"category": "anything"}, "all"))

    def test_missing_website_is_not_tier_a(self):
        score, tier, reason = _score({"name": "Sample", "website": ""})
        self.assertEqual(tier, "B")
        self.assertIn("unverified", reason.lower())
        self.assertLess(score, 60)

    def test_unreachable_site_is_recheck_not_top_priority(self):
        score, tier, reason = _score({
            "name": "Sample", "website": "https://example.org",
            "audit": {"reachable": False}
        })
        self.assertEqual(tier, "C")
        self.assertIn("recheck", reason.lower())

    def test_workbook_has_three_sheets(self):
        columns = [("Business", "name"), ("Score", "score"),
                   ("Status", "website_status"), ("Target market", "target_market")]
        rows = [
            {"name": "Lead A", "score": 40, "website_status": "No website found (unverified)",
             "target_market": "Mehsana, Gujarat, India", "tier": "B"},
            {"name": "Lead B", "score": 10, "website_status": "Functional",
             "target_market": "Mehsana, Gujarat, India", "tier": "C"},
        ]
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "leads.xlsx")
            write_xlsx(rows, columns, path)
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True)
            self.assertEqual(wb.sheetnames, ["Qualified Leads", "All Businesses", "Search Summary"])
            self.assertEqual(wb["All Businesses"].max_row, 3)
            self.assertEqual(wb["Qualified Leads"].max_row, 2)
            wb.close()


if __name__ == "__main__":
    unittest.main()
